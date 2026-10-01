"""One-time seed: turn the two sheet exports into data/knowledge.json.
After this the app never touches Excel. Re-run only if you want to re-seed.

    .venv/bin/python -m backend.seed
"""
import json, re, datetime
from collections import Counter, defaultdict
from pathlib import Path
import openpyxl
from .lookup import norm, first_url, domain, parse_header, KNOWLEDGE

DATA = Path(__file__).resolve().parent.parent / "data"
CLUB_XLSX = DATA / "club_data.xlsx"
PROGRESS_XLSX = DATA / "progress.xlsx"
SKIP = {"prnt.sc", "ibb.co", "skrinshoter.ru", "drive.google.com", "upload-only-hudlvid.s3.amazonaws.com"}


def _rows(ws):
    return [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]


def main():
    cd = openpyxl.load_workbook(CLUB_XLSX, read_only=True, data_only=True)
    teams, tournaments = {}, {}
    for r in _rows(cd["Team Data"])[1:]:
        if r[0]:
            teams[norm(r[0])] = dict(name=str(r[0]).strip(), site=first_url(r[2]), schedule=first_url(r[3]),
                                     mhr=first_url(r[4]), ep=first_url(r[5]), roster=first_url(r[6]), country=None)
    for r in _rows(cd["Club"])[1:]:
        if not r[2]:
            continue
        k = norm(r[2])
        e = dict(name=str(r[2]).strip(), site=first_url(r[4]), schedule=first_url(r[5]), mhr=first_url(r[6]),
                 ep=None, roster=first_url(r[7]), country=str(r[1]).strip() if r[1] else None)
        if k in teams:
            for f in ("site", "schedule", "mhr", "roster", "country"):
                teams[k][f] = teams[k][f] or e[f]
        else:
            teams[k] = e
    for r in _rows(cd["Data"])[1:]:
        if r[2] and r[1] and norm(r[2]) in teams and not teams[norm(r[2])]["country"]:
            teams[norm(r[2])]["country"] = str(r[1]).strip()
    for r in _rows(cd["Tournament data"])[1:]:
        if r[0]:
            tournaments[norm(r[0])] = dict(name=str(r[0]).strip(), site=first_url(r[2]), schedule=first_url(r[3]),
                                           mhr=first_url(r[4]), ep=first_url(r[5]))

    comp_src, team_src = defaultdict(Counter), defaultdict(Counter)
    if PROGRESS_XLSX.exists():
        pw = openpyxl.load_workbook(PROGRESS_XLSX, read_only=True, data_only=True)
        for r in pw["Progress"].iter_rows(values_only=True):
            h = r[2]
            if not (isinstance(h, str) and h.startswith("[")):
                continue
            p = parse_header(h)
            d = domain(first_url(r[8]))
            if p and d and d not in SKIP:
                comp_src[p["compn"]][d] += 1
                team_src[p["t1n"]][d] += 1
                team_src[p["t2n"]][d] += 1

    kb = dict(teams=teams, tournaments=tournaments,
              comp_sources={k: dict(v.most_common(5)) for k, v in comp_src.items()},
              team_sources={k: dict(v.most_common(3)) for k, v in team_src.items()},
              history=[], seeded=datetime.datetime.now().isoformat(timespec="seconds"))
    KNOWLEDGE.write_text(json.dumps(kb, ensure_ascii=False))
    print(f"seeded {len(teams)} teams, {len(tournaments)} tournaments, {len(comp_src)} competition sources -> {KNOWLEDGE}")


if __name__ == "__main__":
    main()
